import argparse
import json
import re
import uuid
from datetime import date
from pathlib import Path

from docx import Document

from app.config import CORPUS_PATH, DEFAULT_SOURCE_DIR
from app.corpus_manifest import build_manifest


REGIME_KEYWORDS = [
    ("WIPO Treaty", "TK"),
    ("Traditional Knowledge Digital Library", "TK"),
    ("TKDL", "TK"),
    ("Biological Diversity", "ABS"),
    ("Nagoya", "ABS"),
    ("Convention on Biological Diversity", "ABS"),
    ("CBD", "ABS"),
    ("Genetic Resources", "ABS"),
    ("Geographical Indications", "GI"),
    ("GI", "GI"),
    ("Trade Marks", "Trademarks"),
    ("Trademark", "Trademarks"),
    ("Madrid", "Trademarks"),
    ("Designs Act", "Designs"),
    ("Design Act", "Designs"),
    ("Hague", "Designs"),
    ("Copyright", "Copyright"),
    ("Plant Varieties", "PlantVariety"),
    ("Farmers Rights", "PlantVariety"),
    ("Drugs and Magic Remedies", "Advertising"),
    ("Ayurveda Aahara", "Food"),
    ("Aahar", "Food"),
    ("FSS", "Food"),
    ("Drugs and Cosmetics", "DrugRegulation"),
    ("Botanical Drug", "DrugRegulation"),
    ("Herbal", "DrugRegulation"),
    ("FDA", "DrugRegulation"),
    ("EMA", "DrugRegulation"),
    ("TRIPS", "Patents"),
    ("Patent Cooperation Treaty", "Patents"),
    ("PCT", "Patents"),
    ("Budapest", "Patents"),
    ("Patent", "Patents"),
]

SOURCE_URLS = [
    ("Patents Act", "https://www.indiacode.nic.in/bitstream/123456789/1392/1/A1970-39.pdf"),
    ("Geographical Indications", "https://www.indiacode.nic.in/bitstream/123456789/1981/1/a199948.pdf"),
    ("Biological Diversity Act", "https://www.indiacode.nic.in/bitstream/123456789/21545/1/the_biological_diversity_act,_2002.pdf"),
    ("Drugs and Cosmetics", "https://www.indiacode.nic.in/bitstream/123456789/18562/1/the_drugs_and_cosmetics_act,_1940.pdf"),
    ("Traditional Knowledge Digital Library", "https://www.csir.res.in/en/documents/tkdl"),
    ("TKDL", "https://www.csir.res.in/en/documents/tkdl"),
    ("TRIPS", "https://www.wto.org/english/docs_e/legal_e/27-trips.pdf"),
    ("Convention on Biological Diversity", "https://www.cbd.int/convention/text/"),
    ("CBD", "https://www.cbd.int/convention/text/"),
    ("Nagoya", "https://www.cbd.int/abs/"),
    ("WIPO Treaty", "https://www.wipo.int/en/web/treaties/ip/gratk/index"),
    ("Patent Cooperation Treaty", "https://www.wipo.int/pct/en/texts/index.html"),
    ("PCT", "https://www.wipo.int/pct/en/texts/index.html"),
    ("Madrid", "https://www.wipo.int/madrid/en/"),
    ("Hague", "https://www.wipo.int/hague/en/"),
    ("Budapest", "https://www.wipo.int/treaties/en/registration/budapest/"),
    ("EMA", "https://www.ema.europa.eu/en/human-regulatory-overview/herbal-medicinal-products"),
    ("FDA", "https://www.fda.gov/regulatory-information/search-fda-guidance-documents/botanical-drug-development-guidance-industry"),
    ("Ayurveda Aahara", "https://fssai.gov.in/"),
]

SECTION_HEADING_RE = re.compile(
    r"^(Section|Article|Rule|Regulation|Schedule|Chapter|Part|Clause|Annex)\s+"
    r"[A-Za-z0-9()./-]+(?:\s*[-:]\s*.+)?$",
    re.IGNORECASE,
)
YEAR_RE = re.compile(r"\b(19|20)\d{2}\b")

SOFT_HEADINGS = {
    "context",
    "purpose",
    "key features",
    "what it is",
    "why it matters",
    "why it was needed",
    "how it works",
    "access and impact",
    "significance",
    "relevance to ip/tk theme",
    "quick recap for exams",
    "note",
    "notes",
}


def clean_text(text: str) -> str:
    return re.sub(r"\s+", " ", text or "").strip()


def strip_leading_number(name: str) -> str:
    return re.sub(r"^\d+[a-zA-Z]?\.\s*", "", name).strip()


def extract_docx_lines(path: Path) -> list[str]:
    doc = Document(path)
    lines = [clean_text(paragraph.text) for paragraph in doc.paragraphs]

    for table in doc.tables:
        for row in table.rows:
            cells = [clean_text(cell.text) for cell in row.cells]
            line = " | ".join(cell for cell in cells if cell)
            lines.append(line)

    return [line for line in lines if line]


def is_probable_heading(line: str) -> bool:
    if not line or len(line) > 120:
        return False
    if SECTION_HEADING_RE.match(line):
        return True
    lowered = line.strip().lower().rstrip(":")
    if lowered in SOFT_HEADINGS:
        return True
    if line.endswith(":") and len(line.split()) <= 8:
        return True
    if line.endswith("."):
        return False
    return len(line.split()) <= 7 and not any(ch in line for ch in ";,")


def is_legal_heading(line: str) -> bool:
    return bool(SECTION_HEADING_RE.match(line or ""))


def infer_source_title(lines: list[str], path: Path) -> str:
    if lines:
        return lines[0]
    return strip_leading_number(path.stem)


def infer_jurisdiction(path: Path, source_title: str) -> str:
    text = f"{path.name} {source_title}".lower()
    international_markers = [
        "trips",
        "convention on biological diversity",
        "nagoya",
        "wipo",
        "patent cooperation treaty",
        "pct",
        "madrid",
        "hague",
        "budapest",
        "fda",
        "ema",
        "international",
    ]
    if any(marker in text for marker in international_markers):
        return "International"
    return "India"


def infer_regime(path: Path, source_title: str, section: str) -> str:
    text = f"{path.name} {source_title} {section}"
    for keyword, regime in REGIME_KEYWORDS:
        if keyword_matches(keyword, text):
            return regime
    return "Patents"


def keyword_matches(keyword: str, text: str) -> bool:
    lowered_keyword = keyword.lower()
    lowered_text = text.lower()
    if len(lowered_keyword) <= 4 and lowered_keyword.isalnum():
        return bool(re.search(rf"\b{re.escape(lowered_keyword)}\b", lowered_text))
    return lowered_keyword in lowered_text


def infer_effective_date(path: Path, source_title: str) -> str:
    text = f"{path.name} {source_title}"
    if "Ayurveda Aahara" in text or "Ayurveda Aahar" in text:
        return "2022-05-05"
    match = YEAR_RE.search(text)
    return match.group(0) if match else ""


def infer_source_url(path: Path, source_title: str) -> str:
    text = f"{path.name} {source_title}"
    for keyword, url in SOURCE_URLS:
        if keyword.lower() in text.lower():
            return url
    if infer_jurisdiction(path, source_title) == "India":
        return "https://www.indiacode.nic.in/"
    return ""


def section_label(heading: str) -> str:
    if not heading or heading == "Overview":
        return "Overview"
    if SECTION_HEADING_RE.match(heading):
        return heading
    return f"Overview - {heading.rstrip(':')}"


def stable_chunk_id(path: Path, heading: str, text: str) -> str:
    namespace = uuid.uuid5(uuid.NAMESPACE_URL, "ip-sakti-corpus")
    key = f"{path.as_posix()}::{heading}::{text[:300]}"
    return str(uuid.uuid5(namespace, key))


def chunk_lines(path: Path, lines: list[str], min_chars: int = 40) -> list[dict]:
    source_title = infer_source_title(lines, path)
    body = lines[1:]
    if body and not is_probable_heading(body[0]) and len(body[0]) <= 120:
        body = body[1:]

    chunks: list[dict] = []
    current_heading = "Overview"
    current_lines: list[str] = []

    def flush() -> None:
        nonlocal current_heading, current_lines
        text = clean_text(" ".join(current_lines))
        if text.strip().lower() == current_heading.strip().lower():
            current_lines = []
            return
        if len(text) < min_chars:
            current_lines = []
            return
        section = section_label(current_heading)
        jurisdiction = infer_jurisdiction(path, source_title)
        chunks.append(
            {
                "chunk_id": stable_chunk_id(path, section, text),
                "source_title": source_title,
                "section_or_article": section,
                "jurisdiction": jurisdiction,
                "regime": infer_regime(path, source_title, section),
                "effective_date": infer_effective_date(path, source_title),
                "doc_version": f"Researcher notes extracted {date.today().isoformat()}",
                "source_url": infer_source_url(path, source_title),
                "source_file": str(path),
                "text": text,
            }
        )
        current_lines = []

    for line in body:
        if is_probable_heading(line):
            if (
                is_legal_heading(current_heading)
                and not is_legal_heading(line)
                and len(current_lines) <= 2
            ):
                current_lines.append(line)
                continue
            flush()
            current_heading = line.rstrip(":")
            current_lines = [line]
        else:
            current_lines.append(line)
    flush()

    if not chunks and lines:
        text = clean_text(" ".join(lines))
        section = "Overview"
        chunks.append(
            {
                "chunk_id": stable_chunk_id(path, section, text),
                "source_title": source_title,
                "section_or_article": section,
                "jurisdiction": infer_jurisdiction(path, source_title),
                "regime": infer_regime(path, source_title, section),
                "effective_date": infer_effective_date(path, source_title),
                "doc_version": f"Researcher notes extracted {date.today().isoformat()}",
                "source_url": infer_source_url(path, source_title),
                "source_file": str(path),
                "text": text,
            }
        )

    return chunks


def build_corpus(
    source_dirs: list[Path] | None = None,
    output_path: Path = CORPUS_PATH,
    min_chars: int = 40,
) -> list[dict]:
    source_dirs = source_dirs or [DEFAULT_SOURCE_DIR]
    all_chunks: list[dict] = []

    for source_dir in source_dirs:
        if not source_dir.exists():
            raise FileNotFoundError(f"source directory not found: {source_dir}")
        for path in sorted(source_dir.glob("*.docx")):
            if path.name.startswith("~$"):
                continue
            lines = extract_docx_lines(path)
            all_chunks.extend(chunk_lines(path, lines, min_chars=min_chars))

    seen: set[str] = set()
    unique_chunks: list[dict] = []
    for chunk in all_chunks:
        if chunk["chunk_id"] in seen:
            continue
        seen.add(chunk["chunk_id"])
        unique_chunks.append(chunk)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(
        json.dumps(unique_chunks, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    output_path.with_name("corpus_manifest.json").write_text(
        json.dumps(build_manifest(unique_chunks), indent=2),
        encoding="utf-8",
    )
    write_changelog(output_path, source_dirs, len(unique_chunks))
    return unique_chunks


def write_changelog(output_path: Path, source_dirs: list[Path], chunk_count: int) -> None:
    changelog = output_path.parent / "corpus_changelog.md"
    source_list = "\n".join(f"- {source_dir}" for source_dir in source_dirs)
    text = (
        "# Corpus Changelog\n\n"
        f"## {date.today().isoformat()}\n"
        f"- Built `{output_path}` from researcher DOCX notes.\n"
        f"- Source directories:\n{source_list}\n"
        f"- Chunk count: {chunk_count}\n"
        "- Chunk IDs are deterministic UUIDv5 values derived from source path, "
        "section label, and chunk text.\n"
    )
    changelog.write_text(text, encoding="utf-8")


def main() -> None:
    parser = argparse.ArgumentParser(description="Build IP-SAKTI corpus JSON from DOCX notes.")
    parser.add_argument(
        "--source-dir",
        action="append",
        type=Path,
        default=None,
        help="Directory containing researcher DOCX notes. Can be repeated.",
    )
    parser.add_argument("--output", type=Path, default=CORPUS_PATH)
    parser.add_argument("--min-chars", type=int, default=40)
    args = parser.parse_args()

    chunks = build_corpus(args.source_dir, args.output, args.min_chars)
    print(f"Wrote {len(chunks)} chunks to {args.output}")


if __name__ == "__main__":
    main()
